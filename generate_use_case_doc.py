import os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell):
    """Set standard 0.5 pt solid black borders on all sides of the cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set inner cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'<w:top w:w="{top}" w:type="dxa"/>\n'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'<w:left w:w="{left}" w:type="dxa"/>\n'
        f'<w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_width(table, width_cm=14.0):
    """Set explicit total table width in dxa."""
    tblPr = table._tbl.tblPr
    width_dxa = int(width_cm * 566.929)  # 1 cm = 566.929 dxa
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tblPr.append(tblW)

def create_use_case_doc():
    doc = Document()

    # Exact Page Setup matching ULM Thesis / KP Ruler:
    # A4 Page: 21.0 cm x 29.7 cm
    # Left Margin: 4.0 cm
    # Right Margin: 3.0 cm
    # Top Margin: 3.0 cm
    # Bottom Margin: 3.0 cm
    # Usable Text Area Width: exactly 14.0 cm (matches 0 to 14 cm on Word ruler!)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(4.0)
        section.right_margin = Cm(3.0)
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(3.0)

    # Base Normal Style: Times New Roman 12pt, 1.15 line spacing, 0 space after
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(18)
    run_t = title_p.add_run("SPESIFIKASI USE CASE SISTEM INTEGRASI DATA BERBASIS WEB")
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(14)
    run_t.font.bold = True

    # 26 Detailed Use Cases
    use_cases = [
        {
            "num": 1,
            "title": "Melihat Status Engine & Metrik Sistem",
            "caption": "Tabel 1 Use case melihat status engine & metrik sistem",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur bagi pengguna untuk melihat kesiapan server backend, status aktifitas konektor database yang terpasang, dan metrik kesehatan sistem pada tampilan dashboard",
            "pre": "Pengguna telah membuka antarmuka aplikasi web Sistem Integrasi Data Berbasis Web",
            "post": "Pengguna berhasil melihat status server aktif dan jumlah konektor yang siap beroperasi",
            "normal": [
                "1. Buka halaman Dashboard",
                "2. Sistem melakukan pengecekan status ke server backend secara otomatis",
                "3. Sistem menampilkan indikator Engine Active beserta daftar konektor yang tersedia (MySQL, PostgreSQL, Oracle, SQLite, CSV, Excel, JSON)"
            ],
            "alt": "-"
        },
        {
            "num": 2,
            "title": "Memilih Template Pipeline Siap Pakai",
            "caption": "Tabel 2 Use case memilih template pipeline siap pakai",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur bagi pengguna untuk memuat konfigurasi alur integrasi data secara instan menggunakan template siap pakai",
            "pre": "Pengguna berada di halaman Dashboard sistem",
            "post": "Pengguna berhasil memuat konfigurasi sumber, pemetaan, dan tujuan ke halaman Pipeline Builder",
            "normal": [
                "1. Pilih kartu template integrasi pada halaman Dashboard",
                "2. Tekan tombol Gunakan Template",
                "3. Sistem memuat preset konfigurasi dan membuka tampilan Pipeline Builder secara otomatis",
                "4. Pengguna dapat langsung meninjau atau menyesuaikan parameter template tersebut"
            ],
            "alt": "-"
        },
        {
            "num": 3,
            "title": "Konfigurasi Sumber Data",
            "caption": "Tabel 3 Use case konfigurasi sumber data",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur induk untuk menentukan dan menghubungkan sumber data yang akan diproses, baik berupa berkas dokumen maupun database relasional enterprise",
            "pre": "Pengguna berada di panel kiri (Source) pada halaman Pipeline Builder",
            "post": "Sumber data terverifikasi dan daftar kolom sumber berhasil dimuat ke antarmuka",
            "normal": [
                "1. Masukkan nama tugas/pipeline pada kolom input nama",
                "2. Pilih kategori sumber data antara File atau Database (SQL)",
                "3. Lengkapi parameter koneksi database atau unggah berkas sumber",
                "4. Sistem memverifikasi akses sumber dan mengaktifkan fitur inspeksi kolom otomatis"
            ],
            "alt": "Apabila parameter koneksi tidak lengkap atau berkas rusak, maka sistem menampilkan pesan kesalahan dan menonaktifkan tombol pemetaan"
        },
        {
            "num": 4,
            "title": "Upload Berkas Dokumen (CSV, Excel, JSON, SQLite)",
            "caption": "Tabel 4 Use case upload berkas dokumen",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk mengunggah satu atau beberapa berkas dokumen lokal yang akan dijadikan sumber data integrasi",
            "pre": "Pengguna memilih kategori sumber File pada Panel Sumber",
            "post": "Pengguna berhasil mengunggah berkas ke server dan struktur kolom siap diinspeksi",
            "normal": [
                "1. Pilih format berkas (CSV, Excel, JSON, JSONL, atau SQLite)",
                "2. Atur karakter delimiter apabila memilih berkas CSV",
                "3. Tarik berkas ke area dropzone atau pilih berkas melalui file explorer",
                "4. Sistem mengunggah berkas ke server dan menampilkan daftar berkas yang berhasil diunggah"
            ],
            "alt": "Apabila format ekstensi berkas tidak didukung atau ukuran berkas melampaui batas server, maka sistem menolak berkas dan menampilkan pesan peringatan format"
        },
        {
            "num": 5,
            "title": "Koneksi Basis Data Relasional (MySQL, PostgreSQL, Oracle)",
            "caption": "Tabel 5 Use case koneksi basis data relasional",
            "actor": "Administrator Sistem, Operator Data",
            "desc": "Fitur untuk mengonfigurasi kredensial koneksi langsung ke server database relasional dan memilih tabel sumber",
            "pre": "Pengguna memilih kategori sumber Database (SQL) pada Panel Sumber",
            "post": "Koneksi database berhasil terhubung dan daftar tabel berhasil dimuat",
            "normal": [
                "1. Pilih jenis database (MySQL, PostgreSQL, atau Oracle)",
                "2. Masukkan parameter host, port, username, password, dan nama database",
                "3. Tekan tombol Tes Koneksi",
                "4. Sistem memverifikasi kredensial ke server database",
                "5. Sistem menampilkan indikator koneksi berhasil dan memunculkan pilihan tabel sumber"
            ],
            "alt": "Apabila host tidak dapat dijangkau atau password salah, maka sistem menampilkan pesan detail error connection refused atau access denied"
        },
        {
            "num": 6,
            "title": "Inspeksi Struktur Kolom Otomatis",
            "caption": "Tabel 6 Use case inspeksi struktur kolom otomatis",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur bagi pengguna untuk melihat seluruh nama kolom serta tipe data bawaan dari sumber data secara otomatis",
            "pre": "Berkas telah diunggah atau tabel database sumber telah dipilih",
            "post": "Daftar seluruh kolom sumber berhasil tampil dalam bentuk checklist interaktif",
            "normal": [
                "1. Sistem melakukan inspeksi metadata kolom secara otomatis dari sumber data terpilih",
                "2. Sistem menampilkan daftar nama kolom sumber beserta kotak centang (checkbox)",
                "3. Pengguna memilih kolom mana saja yang ingin disertakan dalam proses integrasi"
            ],
            "alt": "Apabila berkas kosong atau tabel tidak memiliki kolom, maka sistem menampilkan peringatan bahwa data sumber tidak memiliki kolom yang dapat diproses"
        },
        {
            "num": 7,
            "title": "Override Tipe Kolom Data (Type Casting)",
            "caption": "Tabel 7 Use case override tipe kolom data",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur bagi pengguna untuk mengubah tipe data bawaan kolom sumber secara manual sebelum data diproses dan dimigrasi",
            "pre": "Daftar kolom sumber telah berhasil diinspeksi pada antarmuka",
            "post": "Konfigurasi tipe data kanonikal kolom berhasil disesuaikan dan disimpan",
            "normal": [
                "1. Pengguna melihat daftar kolom sumber yang memiliki dropdown tipe data",
                "2. Pilih kolom yang ingin disesuaikan tipe datanya",
                "3. Pilih tipe data baru melalui dropdown (VARCHAR, INTEGER, FLOAT, DATE, DATETIME)",
                "4. Sistem mencatat aturan casting ini untuk diterapkan secara otomatis saat eksekusi pipeline"
            ],
            "alt": "-"
        },
        {
            "num": 8,
            "title": "Konfigurasi Transformasi dan Pemetaan Kolom",
            "caption": "Tabel 8 Use case konfigurasi transformasi dan pemetaan kolom",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur induk pada Panel Tengah untuk mengelola relasi pemetaan kolom sumber ke tujuan serta menentukan aturan pembersihan data",
            "pre": "Kolom sumber telah dipilih dan jenis tujuan telah ditentukan",
            "post": "Seluruh aturan manipulasi data dan relasi kolom sumber-tujuan tersimpan dengan valid",
            "normal": [
                "1. Buka Panel Tengah (Transform)",
                "2. Atur pasangan pemetaan kolom sumber ke tujuan",
                "3. Tentukan aturan pembersihan data global apabila diperlukan",
                "4. Sistem memvalidasi bahwa tidak ada duplikasi target kolom dan aturan transformasi tersimpan"
            ],
            "alt": "-"
        },
        {
            "num": 9,
            "title": "Pemetaan Kolom Sumber ke Tujuan (1-ke-1 & Ignore)",
            "caption": "Tabel 9 Use case pemetaan kolom sumber ke tujuan",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk menghubungkan kolom sumber ke kolom target yang sesuai, didukung tombol aksi cepat pemetaan massal",
            "pre": "Pengguna berada di bagian Aturan Kolom pada Panel Transformasi",
            "post": "Struktur pemetaan pasangan kolom berhasil tersusun dengan rapi",
            "normal": [
                "1. Periksa daftar kolom sumber di sisi kiri antarmuka pemetaan",
                "2. Tekan tombol Petakan 1-ke-1 untuk memasangkan nama kolom yang identik secara otomatis",
                "3. Pengguna dapat menekan tombol Abaikan Semua untuk mengosongkan relasi kolom",
                "4. Pengguna dapat memilih target kolom tujuan secara manual melalui dropdown",
                "5. Sistem secara otomatis mencegah pemilihan kolom target yang sama dua kali"
            ],
            "alt": "Apabila pengguna memilih kolom target yang sudah digunakan oleh kolom sumber lain, maka sistem mengunci opsi tersebut dan menampilkan peringatan duplikasi"
        },
        {
            "num": 10,
            "title": "Pembersihan & Pengubahan Format Teks",
            "caption": "Tabel 10 Use case pembersihan dan pengubahan format teks",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk menstandarisasi format string teks dengan menghapus spasi liar atau mengubah kapitalisasi huruf",
            "pre": "Pengguna membuka tab Transformasi Global pada Panel Tengah",
            "post": "Aturan normalisasi teks berhasil diaktifkan pada pipeline",
            "normal": [
                "1. Centang opsi Trim Whitespace untuk menghapus spasi di awal dan akhir teks",
                "2. Pilih opsi kapitalisasi teks antara Capitalize, Uppercase, atau Lowercase",
                "3. Sistem memperbarui indikator badge transformasi menjadi aktif"
            ],
            "alt": "-"
        },
        {
            "num": 11,
            "title": "Penanganan Nilai Kosong (Fill Null & Drop Null)",
            "caption": "Tabel 11 Use case penanganan nilai kosong",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk menangani data yang bernilai null atau kosong dengan mengisi nilai pengganti atau membuang baris tersebut",
            "pre": "Pengguna berada di tab Transformasi Global",
            "post": "Aturan sanitasi data kosong berhasil terdaftar dalam parameter eksekusi",
            "normal": [
                "1. Centang opsi Isi Nilai Kosong (Fill Null)",
                "2. Masukkan nilai pengganti default seperti N/A atau 0",
                "3. Pengguna juga dapat memilih opsi Drop Null untuk mengabaikan baris yang memiliki nilai kosong",
                "4. Sistem mencatat aturan penanganan nilai null tersebut"
            ],
            "alt": "-"
        },
        {
            "num": 12,
            "title": "Penyensoran Data Sensitif (Masking)",
            "caption": "Tabel 12 Use case penyensoran data sensitif",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur keamanan untuk menyamarkan (masking) isi kolom yang mengandung data rahasia seperti nomor telepon atau NIK",
            "pre": "Pengguna berada di tab Transformasi Global",
            "post": "Kolom sensitif berhasil disamarkan karakternya saat dimuat ke tujuan",
            "normal": [
                "1. Centang opsi Sensor Data Sensitif (Masking)",
                "2. Masukkan nama-nama kolom yang wajib disamarkan (contoh: phone, nik)",
                "3. Sistem mengonfigurasi modul masking untuk mengganti sebagian karakter menjadi simbol bintang"
            ],
            "alt": "-"
        },
        {
            "num": 13,
            "title": "Penyaringan Baris Data (Row Filtering)",
            "caption": "Tabel 13 Use case penyaringan baris data",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk menyaring baris data yang diproses hanya pada baris yang memenuhi kondisi logika tertentu",
            "pre": "Pengguna berada di tab Transformasi Global",
            "post": "Hanya data yang memenuhi kriteria filter yang akan dimuat ke tujuan",
            "normal": [
                "1. Centang opsi Filter Baris",
                "2. Tentukan nama kolom acuan (contoh: status atau usia)",
                "3. Pilih operator perbandingan (=, !=, >, <, contains)",
                "4. Masukkan nilai ambang batas filter (contoh: ACTIVE atau 18)",
                "5. Sistem mendaftarkan modul filter ke dalam urutan pipeline"
            ],
            "alt": "-"
        },
        {
            "num": 14,
            "title": "Konfigurasi Tujuan Data & Mode Muat",
            "caption": "Tabel 14 Use case konfigurasi tujuan data dan mode muat",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur induk pada Panel Kanan untuk menentukan destinasi penyimpanan data dan strategi pengisian tabel tujuan",
            "pre": "Pengguna berada di Panel Kanan (Destination) pada Pipeline Builder",
            "post": "Parameter target dan mode muat data terkunci siap dieksekusi",
            "normal": [
                "1. Pilih kategori tujuan antara File atau Database (SQL)",
                "2. Lengkapi nama tabel tujuan atau nama berkas ekspor",
                "3. Tentukan salah satu dari 4 mode muat data: Sync, Append, Truncate, atau Replace",
                "4. Sistem memvalidasi kesiapan target penyimpanan"
            ],
            "alt": "-"
        },
        {
            "num": 15,
            "title": "Pemilihan Format Berkas Ekspor Target",
            "caption": "Tabel 15 Use case pemilihan format berkas ekspor target",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk mengekspor seluruh hasil integrasi data ke dalam berkas fisik yang dapat diunduh",
            "pre": "Pengguna memilih kategori tujuan File pada Panel Destination",
            "post": "Sistem siap menuliskan aliran data ke dalam berkas fisik server",
            "normal": [
                "1. Pilih format berkas keluaran: CSV, Excel (.xlsx), SQLite (.db), JSON, atau JSONL",
                "2. Tentukan nama berkas keluaran yang diinginkan",
                "3. Sistem menetapkan ekstensi berkas secara otomatis sesuai format terpilih"
            ],
            "alt": "-"
        },
        {
            "num": 16,
            "title": "Pemilihan Basis Data Target (MySQL, PostgreSQL, Oracle)",
            "caption": "Tabel 16 Use case pemilihan basis data target",
            "actor": "Administrator Sistem, Operator Data",
            "desc": "Fitur untuk memuat hasil integrasi data langsung ke dalam server database relasional target",
            "pre": "Pengguna memilih kategori tujuan Database (SQL) pada Panel Destination",
            "post": "Koneksi database tujuan terkonfigurasi dan nama tabel target terdaftar",
            "normal": [
                "1. Pilih jenis database tujuan (MySQL, PostgreSQL, atau Oracle)",
                "2. Masukkan parameter host, port, username, password, dan nama database tujuan",
                "3. Tentukan nama tabel target tempat data akan disimpan"
            ],
            "alt": "-"
        },
        {
            "num": 17,
            "title": "Pemilihan Mode Muat (Sync, Append, Truncate, Replace)",
            "caption": "Tabel 17 Use case pemilihan mode muat data",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk menentukan strategi perlakuan terhadap data lama pada tabel tujuan saat proses pemindahan data berlangsung",
            "pre": "Target database tujuan telah ditentukan",
            "post": "Mode pengisian data terpilih dan siap dijalankan oleh engine",
            "normal": [
                "1. Pilih salah satu opsi mode pengisian:",
                "   a. Sync / Mirror: Menyelaraskan data secara utuh beserta evolusi skema penambahan dan penghapusan kolom",
                "   b. Append: Menyisipkan baris baru di bawah data lama tanpa menghapus data sebelumnya",
                "   c. Truncate & Insert: Mengosongkan data lama terlebih dahulu, lalu memasukkan data baru",
                "   d. Replace Table: Menghapus tabel lama (DROP) dan membuat struktur tabel baru (CREATE)",
                "2. Sistem mengunci mode pemuatan yang dipilih"
            ],
            "alt": "-"
        },
        {
            "num": 18,
            "title": "Konfirmasi Peringatan Operasi Destruktif",
            "caption": "Tabel 18 Use case konfirmasi peringatan operasi destruktif",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur dialog keamanan sistem untuk mencegah penghapusan data secara tidak sengaja saat pengguna memilih mode Replace atau Truncate",
            "pre": "Pengguna memilih mode Replace atau Truncate dan menekan tombol Jalankan Pipeline",
            "post": "Pengguna mengonfirmasi persetujuan risiko atau membatalkan operasi destruktif",
            "normal": [
                "1. Sistem memunculkan Modal Peringatan Bahaya dengan warna aksen merah",
                "2. Sistem menampilkan pesan bahwa seluruh data lama pada tabel target akan dihapus permanen",
                "3. Pengguna membaca peringatan dan menekan tombol Saya Mengerti, Lanjutkan",
                "4. Sistem melanjutkan proses eksekusi pipeline"
            ],
            "alt": "Apabila pengguna menekan tombol Batal, maka sistem menutup modal dialog dan membatalkan proses eksekusi"
        },
        {
            "num": 19,
            "title": "Eksekusi Pipeline Data",
            "caption": "Tabel 19 Use case eksekusi pipeline data",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur utama untuk memicu eksekusi pemrosesan data secara menyeluruh dari ekstraksi, transformasi, hingga pemuatan ke tujuan",
            "pre": "Seluruh parameter pada Panel Sumber, Transformasi, dan Tujuan telah terisi lengkap",
            "post": "Data berhasil dimigrasikan ke target dan riwayat tercatat di sistem",
            "normal": [
                "1. Tekan tombol Jalankan Pipeline",
                "2. Frontend mengirimkan request ke backend runner",
                "3. Backend menginisialisasi eksekusi di latar belakang agar server tidak terblokir",
                "4. Sistem membuka Progress Drawer pemantauan real-time secara otomatis"
            ],
            "alt": "Apabila parameter wajib belum diisi (misal nama tabel target kosong), maka sistem menampilkan pesan peringatan validasi dan membatalkan eksekusi"
        },
        {
            "num": 20,
            "title": "Verifikasi Evolusi Skema (Schema Evolution)",
            "caption": "Tabel 20 Use case verifikasi evolusi skema",
            "actor": "Sistem (Otomatis saat eksekusi pipeline)",
            "desc": "Fitur kecerdasan backend untuk mendeteksi perubahan struktur kolom antara sumber dan target serta melakukan penyesuaian otomatis",
            "pre": "Pipeline mulai berjalan menuju database target",
            "post": "Struktur kolom tabel target tersinkronisasi dengan struktur kolom sumber data",
            "normal": [
                "1. Backend menginspeksi apakah tabel target sudah ada di database tujuan",
                "2. Apabila belum ada, sistem mengeksekusi DDL CREATE TABLE baru",
                "3. Apabila tabel target sudah ada, sistem membandingkan daftar kolom sumber dengan target",
                "4. Apabila ditemukan kolom baru pada sumber, sistem mengeksekusi ALTER TABLE ADD COLUMN",
                "5. Apabila mode Sync dipilih dan ada kolom yang hilang pada sumber, sistem mengeksekusi ALTER TABLE DROP COLUMN"
            ],
            "alt": "-"
        },
        {
            "num": 21,
            "title": "Pemantauan Progres Real-Time (SSE Stream)",
            "caption": "Tabel 21 Use case pemantauan progres real-time",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk mengamati jalannya proses integrasi data detik demi detik melalui panel drawer geser tanpa perlu memuat ulang halaman web",
            "pre": "Pipeline data sedang berjalan di background server",
            "post": "Pengguna berhasil melihat transparansi status proses, persentase progres, dan log langkah teknis",
            "normal": [
                "1. Progress Drawer terbuka di sisi kanan antarmuka",
                "2. Frontend membuat koneksi streaming Server-Sent Events (SSE) ke backend",
                "3. Progress bar bergerak halus dari 0% hingga 100% mengikuti event data",
                "4. Indikator langkah kerja (Connect -> Extract -> Transform -> Load) menyala hijau secara berurutan",
                "5. Area terminal log monospace menampilkan catatan baris demi baris pemrosesan",
                "6. Saat selesai 100%, sistem menampilkan ringkasan total baris berhasil dimuat dan durasi waktu"
            ],
            "alt": "Apabila terjadi kegagalan di tengah proses, maka sistem menangkap error, mengirim status error via SSE, menampilkan pesan traceback yang informatif, dan menyediakan tombol Coba Lagi"
        },
        {
            "num": 22,
            "title": "Unduh Berkas Hasil Ekspor Langsung dari Drawer",
            "caption": "Tabel 22 Use case unduh berkas hasil ekspor langsung dari drawer",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur untuk mengunduh langsung berkas fisik hasil ekspor begitu proses pipeline di drawer selesai 100%",
            "pre": "Eksekusi pipeline selesai sukses dan target keluaran berupa berkas",
            "post": "Berkas hasil berhasil terunduh ke komputer lokal pengguna",
            "normal": [
                "1. Pipeline menyelesaikan 100% proses dan menghasilkan berkas keluaran",
                "2. Sistem memunculkan tombol Unduh Berkas Hasil pada bagian bawah Progress Drawer",
                "3. Tekan tombol unduh berkas",
                "4. Browser mengunduh berkas fisik ke direktori download perangkat pengguna"
            ],
            "alt": "-"
        },
        {
            "num": 23,
            "title": "Kelola Penjadwalan Tugas (Scheduler)",
            "caption": "Tabel 23 Use case kelola penjadwalan tugas",
            "actor": "Administrator Sistem",
            "desc": "Fitur untuk mengelola daftar otomatisasi pipeline yang berjalan secara berkala menggunakan cron di latar belakang",
            "pre": "Pengguna membuka menu Penjadwalan (Schedules)",
            "post": "Tabel daftar jadwal aktif dan nonaktif tampil lengkap beserta kontrol pengelolaannya",
            "normal": [
                "1. Buka halaman Penjadwalan melalui navigasi",
                "2. Sistem memuat daftar seluruh jadwal otomatis dari server",
                "3. Sistem menampilkan tabel jadwal yang mencakup nama, frekuensi, status, waktu terakhir jalan, dan jadwal berikutnya",
                "4. Pengguna dapat mengubah status aktif atau nonaktif melalui switch toggle",
                "5. Pengguna dapat menghapus jadwal yang sudah tidak digunakan"
            ],
            "alt": "-"
        },
        {
            "num": 24,
            "title": "Tambah Jadwal Cron Baru & Picu Eksekusi Manual",
            "caption": "Tabel 24 Use case tambah jadwal cron baru dan picu eksekusi manual",
            "actor": "Administrator Sistem",
            "desc": "Fitur untuk mendaftarkan jadwal otomatisasi baru melalui formulir modal dan memicu eksekusi seketika tanpa menunggu waktu cron",
            "pre": "Pengguna berada di halaman Penjadwalan",
            "post": "Jadwal baru berhasil tersimpan di sistem atau eksekusi seketika berhasil dipicu",
            "normal": [
                "1. Tekan tombol Tambah Jadwal Baru",
                "2. Sistem membuka Modal Form Penjadwalan",
                "3. Masukkan nama jadwal dan tentukan frekuensi waktu eksekusi",
                "4. Tekan tombol Simpan Jadwal untuk mendaftarkan jadwal ke database",
                "5. Pengguna juga dapat menekan tombol Jalankan Sekarang pada baris tabel untuk memicu eksekusi manual seketika"
            ],
            "alt": "-"
        },
        {
            "num": 25,
            "title": "Lihat Riwayat Eksekusi & Audit Log Error",
            "caption": "Tabel 25 Use case lihat riwayat eksekusi dan audit log error",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur audit rekam jejak untuk memeriksa histori seluruh proses pemindahan data yang pernah dijalankan sebelumnya",
            "pre": "Pengguna membuka menu History",
            "post": "Tabel riwayat eksekusi lampau berhasil ditampilkan lengkap",
            "normal": [
                "1. Buka halaman History melalui navigasi",
                "2. Sistem menampilkan tabel daftar eksekusi yang memuat ID Run, nama pipeline, status keberhasilan, waktu eksekusi, durasi detik, dan total baris data",
                "3. Pengguna dapat mengklik baris riwayat untuk melihat rincian pesan error atau catatan log proses teknis"
            ],
            "alt": "-"
        },
        {
            "num": 26,
            "title": "Kelola & Unduh Berkas Data Hasil Ekspor (File Manager)",
            "caption": "Tabel 26 Use case kelola dan unduh berkas data hasil ekspor",
            "actor": "Operator Data, Administrator Sistem",
            "desc": "Fitur manajemen berkas fisik di server untuk melihat, mengunduh ulang, atau menghapus file hasil keluaran integrasi",
            "pre": "Pengguna membuka menu File Data (Unduh)",
            "post": "Pengguna berhasil mengunduh berkas fisik yang tersimpan di server",
            "normal": [
                "1. Buka halaman File Data melalui navigasi",
                "2. Sistem menampilkan tabel daftar berkas yang mencakup nama berkas, format ekstensi, ukuran berkas (KB/MB), dan tanggal pembuatan",
                "3. Tekan tombol Unduh pada berkas yang diinginkan",
                "4. Browser mengunduh berkas fisik ke perangkat lokal pengguna"
            ],
            "alt": "Apabila berkas fisik telah terhapus dari server penyimpanan, maka sistem menampilkan pesan peringatan bahwa berkas sudah tidak tersedia"
        }
    ]

    for uc in use_cases:
        # Title paragraph (e.g. "15.  Cek Status Retraining")
        p_num = doc.add_paragraph()
        p_num.paragraph_format.space_before = Pt(14)
        p_num.paragraph_format.space_after = Pt(4)
        p_num.paragraph_format.line_spacing = 1.15
        r_num = p_num.add_run(f"{uc['num']}.  {uc['title']}")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(12)
        # Note: In screenshot, heading is regular or bold Times New Roman
        r_num.font.bold = False

        # Caption paragraph (e.g. "Tabel 16 Use case cek status retraining")
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.line_spacing = 1.15
        r_cap = p_cap.add_run(uc['caption'])
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(12)

        # Create Table (7 rows, 2 columns)
        table = doc.add_table(rows=7, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Exact widths matching the Word ruler in screenshot:
        # Ruler 0 to 14.0 cm: Column 1 = 3.6 cm, Column 2 = 10.4 cm. Total = 14.0 cm
        set_table_width(table, width_cm=14.0)
        col_widths = [Cm(3.6), Cm(10.4)]
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

        # Row 0: Header ("Elemen" | "Deskripsi")
        hdr_cells = table.rows[0].cells
        
        p_h1 = hdr_cells[0].paragraphs[0]
        p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h1.paragraph_format.space_before = Pt(0)
        p_h1.paragraph_format.space_after = Pt(0)
        p_h1.paragraph_format.line_spacing = 1.15
        r_h1 = p_h1.add_run("Elemen")
        r_h1.font.name = 'Times New Roman'
        r_h1.font.size = Pt(12)
        r_h1.font.bold = True

        p_h2 = hdr_cells[1].paragraphs[0]
        p_h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h2.paragraph_format.space_before = Pt(0)
        p_h2.paragraph_format.space_after = Pt(0)
        p_h2.paragraph_format.line_spacing = 1.15
        r_h2 = p_h2.add_run("Deskripsi")
        r_h2.font.name = 'Times New Roman'
        r_h2.font.size = Pt(12)
        r_h2.font.bold = True

        data_rows = [
            ("Aktor Utama", uc['actor']),
            ("Deskripsi", uc['desc']),
            ("Precondition", uc['pre']),
            ("Postcondition", uc['post']),
            ("Alur Normal", "\n".join(uc['normal'])),
            ("Alur Alternatif", uc['alt'])
        ]

        for i, (elem, desc) in enumerate(data_rows, start=1):
            row_cells = table.rows[i].cells
            
            # Left cell: Elemen (Regular font, matching screenshot!)
            p_e = row_cells[0].paragraphs[0]
            p_e.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_e.paragraph_format.space_before = Pt(0)
            p_e.paragraph_format.space_after = Pt(0)
            p_e.paragraph_format.line_spacing = 1.15
            r_e = p_e.add_run(elem)
            r_e.font.name = 'Times New Roman'
            r_e.font.size = Pt(12)
            r_e.font.bold = False

            # Right cell: Deskripsi
            p_d = row_cells[1].paragraphs[0]
            p_d.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_d.paragraph_format.space_before = Pt(0)
            p_d.paragraph_format.space_after = Pt(0)
            p_d.paragraph_format.line_spacing = 1.15
            r_d = p_d.add_run(desc)
            r_d.font.name = 'Times New Roman'
            r_d.font.size = Pt(12)
            r_d.font.bold = False

        # Apply borders, padding, and alignments to all cells
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_border(cell)
                # Cell padding: top/bottom ~5pt (100 dxa), left/right ~7pt (140 dxa)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

        # Space between tables
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(10)

    output_path = "/Users/ghn/.gemini/antigravity/scratch/ngoding anjay/TABEL_USE_CASE_SISTEM_INTEGRASI_DATA.docx"
    doc.save(output_path)
    print(f"File berhasil diupdate dengan ukuran presisi di: {output_path}")

if __name__ == "__main__":
    create_use_case_doc()
