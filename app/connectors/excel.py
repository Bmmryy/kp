"""FlowETL Excel (XLSX) Source and Destination Connectors.

Provides spreadsheet extraction and loading for Microsoft Excel (.xlsx) files
using openpyxl, featuring formatted headers, grid auto-filter, and column width auto-sizing.
"""

from __future__ import annotations

import datetime
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.connectors.base import ConnectorConfig, DestinationConnector, SourceConnector
from app.core.exceptions import (
    ConnectionFailedError,
    ExtractionError,
    LoadError,
    SchemaDiscoveryError,
)
from app.schema.models import ColumnDefinition, Dataset, DataType, TableSchema
from app.utils.logging import get_logger

logger = get_logger("flowetl.connectors.excel")


class ExcelSource(SourceConnector):
    """Source connector for reading tables from Excel (.xlsx) spreadsheets."""

    def __init__(self, config: ConnectorConfig) -> None:
        super().__init__(config)
        self.base_path = Path(config.options.get("path", "."))
        self.sheet_name = config.options.get("sheet_name")

    def _resolve_file_path(self, table_name: str) -> Path:
        if self.base_path.suffix or (self.base_path.exists() and self.base_path.is_file()):
            return self.base_path
        target = self.base_path / table_name
        if not target.suffix:
            target = target.with_suffix(".xlsx")
        return target

    def connect(self) -> None:
        if not self.base_path.exists():
            raise ConnectionFailedError(
                message=f"Excel source path does not exist: {self.base_path.resolve()}",
                suggested_action="Verify the 'path' parameter in your Excel connection configuration.",
            )
        self._is_connected = True

    def test_connection(self) -> bool:
        self.connect()
        return True

    def list_tables(self) -> List[str]:
        if not self.is_connected:
            self.connect()
        if self.base_path.is_file():
            wb = openpyxl.load_workbook(self.base_path, read_only=True)
            sheets = wb.sheetnames
            wb.close()
            return sheets
        return sorted([f.stem for f in self.base_path.glob("*.xlsx")])

    def get_schema(self, table_name: str) -> TableSchema:
        if not self.is_connected:
            self.connect()
        file_path = self._resolve_file_path(table_name)
        if not file_path.exists():
            raise SchemaDiscoveryError(
                message=f"Excel file not found for table '{table_name}' at {file_path.resolve()}.",
                suggested_action="Verify table/file name or check directory contents.",
            )

        wb = openpyxl.load_workbook(file_path, read_only=True)
        sheet_to_use = self.sheet_name or (table_name if table_name in wb.sheetnames else wb.sheetnames[0])
        ws = wb[sheet_to_use]

        header_row = next(ws.iter_rows(values_only=True), None)
        wb.close()

        if not header_row:
            return TableSchema(name=table_name, columns=[])

        columns: List[ColumnDefinition] = []
        for col_name in header_row:
            if col_name is not None and str(col_name).strip():
                columns.append(
                    ColumnDefinition(
                        name=str(col_name).strip(),
                        data_type=DataType.STRING,
                        nullable=True,
                    )
                )

        return TableSchema(name=table_name, columns=columns)

    def extract(
        self,
        table_name: str,
        chunk_size: int = 1000,
        query: Optional[str] = None,
    ) -> Iterator[Dataset]:
        if not self.is_connected:
            self.connect()
        file_path = self._resolve_file_path(table_name)
        schema = self.get_schema(table_name)

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet_to_use = self.sheet_name or (table_name if table_name in wb.sheetnames else wb.sheetnames[0])
        ws = wb[sheet_to_use]

        rows_iter = ws.iter_rows(values_only=True)
        headers = next(rows_iter, None)
        if not headers:
            wb.close()
            yield Dataset(schema=schema, rows=[])
            return

        col_keys = [str(h).strip() for h in headers if h is not None]

        batch: List[Dict[str, Any]] = []
        for raw_row in rows_iter:
            row_dict = {}
            for i, k in enumerate(col_keys):
                val = raw_row[i] if i < len(raw_row) else None
                row_dict[k] = val
            batch.append(row_dict)
            if len(batch) >= chunk_size:
                yield Dataset(schema=schema, rows=batch)
                batch = []

        if batch:
            yield Dataset(schema=schema, rows=batch)
        wb.close()

    def close(self) -> None:
        self._is_connected = False


class ExcelDestination(DestinationConnector):
    """Destination connector for writing datasets into styled Excel (.xlsx) spreadsheets."""

    def __init__(self, config: ConnectorConfig) -> None:
        super().__init__(config)
        self.base_path = Path(config.options.get("path", "."))
        self.sheet_name = config.options.get("sheet_name")

    def _resolve_file_path(self, table_name: str) -> Path:
        if self.base_path.suffix or (self.base_path.exists() and self.base_path.is_file()):
            return self.base_path
        target = self.base_path / table_name
        if not target.suffix:
            target = target.with_suffix(".xlsx")
        return target

    def connect(self) -> None:
        if self.base_path.suffix or (self.base_path.exists() and self.base_path.is_file()):
            self.base_path.parent.mkdir(parents=True, exist_ok=True)
        else:
            self.base_path.mkdir(parents=True, exist_ok=True)
        self._is_connected = True

    def test_connection(self) -> bool:
        self.connect()
        return True

    @staticmethod
    def _apply_header_style(ws, headers: List[str]) -> None:
        header_fill = PatternFill(start_color="1877F2", end_color="1877F2", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )
        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"

    def create_schema(self, schema: TableSchema, if_exists: str = "fail") -> None:
        if not self.is_connected:
            self.connect()
        file_path = self._resolve_file_path(schema.name)
        if file_path.exists() and if_exists == "fail":
            raise LoadError(
                message=f"Excel destination file '{file_path.name}' already exists.",
                suggested_action="Set load mode to 'replace' or specify a different file name.",
            )

        if not file_path.exists() or if_exists in ("replace", "truncate", "sync", "mirror"):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = (self.sheet_name or schema.name or "Sheet1")[:31]
            if schema.column_names:
                self._apply_header_style(ws, schema.column_names)
            wb.save(file_path)
            wb.close()

    def load(
        self,
        dataset: Dataset,
        table_name: str,
        mode: str = "append",
    ) -> int:
        if not self.is_connected:
            self.connect()
        file_path = self._resolve_file_path(table_name)
        file_exists = file_path.exists()

        if mode in ("replace", "truncate", "sync", "mirror") or not file_exists:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = (self.sheet_name or table_name or "Sheet1")[:31]
            self._apply_header_style(ws, dataset.schema_def.column_names)
            start_row = 2
        else:
            wb = openpyxl.load_workbook(file_path)
            title = (self.sheet_name or table_name or "Sheet1")[:31]
            if title in wb.sheetnames:
                ws = wb[title]
            else:
                ws = wb.create_sheet(title=title)

            # Check if row 1 has headers
            has_headers = ws.max_row >= 1 and any(ws.cell(row=1, column=c).value is not None for c in range(1, len(dataset.schema_def.column_names) + 1))
            if not has_headers:
                self._apply_header_style(ws, dataset.schema_def.column_names)
                start_row = 2
            else:
                start_row = ws.max_row + 1

        headers = dataset.schema_def.column_names
        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )

        curr_row = start_row
        for r_dict in dataset.rows:
            for col_idx, col_name in enumerate(headers, 1):
                raw_val = r_dict.get(col_name)
                if isinstance(raw_val, datetime.timedelta):
                    total_sec = int(raw_val.total_seconds())
                    h, rem = divmod(total_sec, 3600)
                    m, s = divmod(rem, 60)
                    val = f"{h:02}:{m:02}:{s:02}"
                else:
                    val = raw_val

                cell = ws.cell(row=curr_row, column=col_idx, value=val)
                cell.border = thin_border
                if curr_row % 2 == 0:
                    cell.fill = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
            curr_row += 1

        # Auto-fit column widths
        for col_idx, col_name in enumerate(headers, 1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(col_name))
            for r in range(1, min(ws.max_row + 1, 100)):
                cv = ws.cell(row=r, column=col_idx).value
                if cv is not None:
                    max_len = max(max_len, len(str(cv)[:40]))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        # Enable auto-filter
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"

        wb.save(file_path)
        wb.close()
        return dataset.row_count

    def close(self) -> None:
        self._is_connected = False
